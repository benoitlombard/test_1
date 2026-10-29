# Builds the results workbook. Separate from main.py so the pipeline logic
# and the report formatting don't get tangled together.

import os###################
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREEN = ("C6EFCE", "006100")
YELLOW = ("FFEB9C", "9C6500")
RED = ("FFC7CE", "9C0006")
STATUS_COLORS = {"GO": GREEN, "REVIEW": YELLOW, "NO-GO": RED}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BOLD = Font(name="Arial", size=10, bold=True)
BODY = Font(name="Arial", size=10)
BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)

TABLE_HEADERS = ["Candidate", "Score", "Missing", "Contact", "Mandatory met"]
HEADER_ROW = 7


def save_result_xlsx(rfq_text, result, rows, out_dir="results"):
    os.makedirs(out_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(out_dir, f"result_{timestamp}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "RFQ Match Result"

    ws["A1"] = "RFQ Match Result"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws.merge_cells("A1:E1")

    ws["A3"] = "RFQ"
    ws["A3"].font = BOLD
    ws["B3"] = rfq_text.strip()
    ws["B3"].font = BODY
    ws["B3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B3:E3")

    fill, color = STATUS_COLORS[result["status"]]
    ws["A4"] = "Result"
    ws["A4"].font = BOLD
    ws["B4"] = result["status"]
    ws["B4"].font = Font(name="Arial", size=11, bold=True, color=color)
    ws["B4"].fill = PatternFill("solid", fgColor=fill)
    ws["B4"].alignment = Alignment(horizontal="center")

    ws["A5"] = "Recommended"
    ws["A5"].font = BOLD
    top = result["top_candidate"]
    if top:
        ws["B5"] = f"{top['profile']['name']} ({top['profile'].get('contact', 'no contact on file')})"
    else:
        ws["B5"] = "No candidate meets every mandatory requirement"
    ws["B5"].font = BODY

    for col, title in enumerate(TABLE_HEADERS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER

    for i, row in enumerate(rows):
        r = HEADER_ROW + 1 + i
        values = [row["name"], row["score"], row["missing"], row["contact"], "Yes" if row["mandatory_met"] else "No"]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = BODY
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if col != 3 else "left")
        ws.cell(row=r, column=2).number_format = '0"%"'

        fill, color = GREEN if row["mandatory_met"] else RED
        met_cell = ws.cell(row=r, column=5)
        met_cell.fill = PatternFill("solid", fgColor=fill)
        met_cell.font = Font(name="Arial", size=10, color=color)

    last_row = HEADER_ROW + len(rows)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(TABLE_HEADERS))}{last_row}"
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    for col, width in enumerate([18, 10, 32, 28, 16], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(path)
    return path
